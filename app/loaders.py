import io
import json
import re
from dataclasses import dataclass
from html.parser import HTMLParser

from pypdf import PdfReader

TEXT_EXTENSIONS = {"txt", "md", "markdown", "csv", "tsv", "log", "rst", "xml", "yaml", "yml"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | {"pdf", "docx", "html", "htm", "json"}


class UnsupportedFileError(ValueError):
    pass


@dataclass
class Page:
    text: str
    number: int | None = None


def extension(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _normalize(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


class _HTMLText(HTMLParser):
    _skip = {"script", "style", "noscript", "svg", "head"}
    _block = {
        "p",
        "div",
        "br",
        "li",
        "tr",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "section",
        "article",
    }

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._depth = 0

    def handle_starttag(self, tag, attrs):
        if tag in self._skip:
            self._depth += 1
        elif tag in self._block:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self._skip and self._depth:
            self._depth -= 1
        elif tag in self._block:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self._depth:
            self.parts.append(data)


def _pdf(data: bytes) -> list[Page]:
    reader = PdfReader(io.BytesIO(data))
    return [Page(page.extract_text() or "", i + 1) for i, page in enumerate(reader.pages)]


def _docx(data: bytes) -> list[Page]:
    from docx import Document

    doc = Document(io.BytesIO(data))
    blocks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return [Page("\n".join(blocks))]


def _html(data: bytes) -> list[Page]:
    parser = _HTMLText()
    parser.feed(_decode(data))
    return [Page("".join(parser.parts))]


def _json(data: bytes) -> list[Page]:
    raw = _decode(data)
    try:
        return [Page(json.dumps(json.loads(raw), indent=2, ensure_ascii=False))]
    except json.JSONDecodeError:
        return [Page(raw)]


def parse(filename: str, data: bytes) -> list[Page]:
    ext = extension(filename)
    if ext == "pdf":
        pages = _pdf(data)
    elif ext == "docx":
        pages = _docx(data)
    elif ext in {"html", "htm"}:
        pages = _html(data)
    elif ext == "json":
        pages = _json(data)
    elif ext in TEXT_EXTENSIONS:
        pages = [Page(_decode(data))]
    else:
        raise UnsupportedFileError(f"Unsupported file type: .{ext or '?'}")
    return [Page(_normalize(p.text), p.number) for p in pages if p.text and p.text.strip()]
